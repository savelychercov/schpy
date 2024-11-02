import sys
from PyQt5.QtWidgets import QApplication, QMainWindow, QTableWidget, QTableWidgetItem, QPushButton, QVBoxLayout, \
    QHBoxLayout, QWidget, QAbstractItemView, QMessageBox, QLabel, QListWidget, QDialog, QCheckBox
from PyQt5.QtCore import Qt, QEvent
from PyQt5.QtGui import QIcon
import schedule_maker
import db
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from datetime import time

data: db.Data


class InputDataDialog(QDialog):
    def __init__(self):
        super().__init__()

        with open("InputDataDialog.css", "r") as f:
            self.setStyleSheet(f.read())

        self.vars_to_redact = {
            "Смены групп": "groups_shift",
            "КУГ": "discipline_hours",
            "Преподаватели": "teachers",
            "Аудитории": "rooms",
            "Расписание\nпреподавателей": "teachers_work_hours",
            "Расписание\nаудиторий": "rooms_availability_hours",
        }

        self.setWindowTitle("Редактирование данных")
        self.setGeometry(100, 100, 1200, 600)
        self.setWindowFlags(self.windowFlags() | Qt.WindowStaysOnTopHint)

        self.layout = QHBoxLayout(self)

        self.variable_list = QListWidget()
        self.variable_list.setFixedWidth(250)
        self.variable_list.addItems(list(self.vars_to_redact.keys()))
        self.variable_list.currentItemChanged.connect(self.display_variable_data)
        self.layout.addWidget(self.variable_list)

        self.data_table = QTableWidget()
        self.layout.addWidget(self.data_table)

        self.button_layout = QVBoxLayout()

        self.add_row_button = QPushButton("Добавить строку")
        self.add_row_button.clicked.connect(self.add_row)
        self.button_layout.addWidget(self.add_row_button)

        self.delete_selected_rows_button = QPushButton("Удалить выделенные строки")
        self.delete_selected_rows_button.clicked.connect(self.delete_selected_rows)
        self.button_layout.addWidget(self.delete_selected_rows_button)

        self.save_button = QPushButton("Сохранить изменения")
        self.save_button.clicked.connect(self.save_changes)
        self.button_layout.addWidget(self.save_button)

        self.back_button = QPushButton("Назад")
        self.back_button.clicked.connect(self.close)
        self.button_layout.addWidget(self.back_button)

        self.button_layout.addStretch()
        self.layout.addLayout(self.button_layout)

        self.current_variable = None

    def delete_selected_rows(self):
        resp = QMessageBox.question(
            self, "Подтверждение", "Удалить выделенные строки?", QMessageBox.Yes | QMessageBox.No
        )
        if resp == QMessageBox.Yes:
            selected_rows = self.data_table.selectionModel().selectedRows()
            for row in selected_rows:
                self.data_table.removeRow(row.row())

    def add_row(self):
        self.data_table.insertRow(self.data_table.rowCount())

    def display_variable_data(self, current):
        if not current:
            return

        variable_name = self.vars_to_redact[current.text()]
        self.current_variable = variable_name
        variable_data = getattr(data, variable_name)

        self.data_table.setRowCount(0)

        if variable_name == "groups_shift":
            headers = ["Группа", "Пара", "Время", "Тип"]
            self.data_table.setColumnCount(len(headers))
            self.data_table.setHorizontalHeaderLabels(headers)
            row = 0
            for group, schedule in variable_data.items():
                for pair, time_obj in schedule.items():
                    self.data_table.insertRow(row)
                    self.data_table.setItem(row, 0, QTableWidgetItem(group))
                    self.data_table.setItem(row, 1, QTableWidgetItem(str(pair)))
                    self.data_table.setItem(row, 2, QTableWidgetItem(time_obj.get_str()))
                    self.data_table.setItem(row, 3, QTableWidgetItem(time_obj.pair_type))
                    row += 1

        elif variable_name == "discipline_hours":
            self.data_table.setColumnCount(3)
            self.data_table.setHorizontalHeaderLabels(["Группа", "Дисциплина", "Часы"])
            row = 0
            for group, disciplines in variable_data.items():
                for discipline, hours in disciplines.items():
                    self.data_table.insertRow(row)
                    self.data_table.setItem(row, 0, QTableWidgetItem(group))
                    self.data_table.setItem(row, 1, QTableWidgetItem(discipline))
                    self.data_table.setItem(row, 2, QTableWidgetItem(str(hours)))
                    row += 1

        elif variable_name == "teachers":
            self.data_table.setColumnCount(3)
            self.data_table.setHorizontalHeaderLabels(["Имя", "Дисциплины", "Группы"])
            row = 0
            for name, teacher in variable_data.items():
                self.data_table.insertRow(row)
                self.data_table.setItem(row, 0, QTableWidgetItem(name))
                self.data_table.setItem(row, 1, QTableWidgetItem(", ".join(teacher.disciplines)))
                self.data_table.setItem(row, 2, QTableWidgetItem(", ".join(teacher.groups)))
                row += 1

        elif variable_name == "rooms":
            self.data_table.setColumnCount(2)
            self.data_table.setHorizontalHeaderLabels(["Аудитория", "Онлайн"])
            row = 0
            for room, room_obj in variable_data.items():
                self.data_table.insertRow(row)
                self.data_table.setItem(row, 0, QTableWidgetItem(room))
                self.data_table.setItem(row, 1, QTableWidgetItem("Да" if room_obj.is_online else "Нет"))
                row += 1

        elif variable_name in ["teachers_work_hours", "rooms_availability_hours"]:
            self._setup_schedule_table(variable_data)

        self.data_table.setEditTriggers(QTableWidget.DoubleClicked)
        self.data_table.resizeColumnsToContents()

    def _setup_schedule_table(self, schedule_data):
        days_of_week = list(data.days)
        num_pairs = len(data.teachers_schedule_time)

        # Устанавливаем столбцы: один для имени, остальные для расписания по дням
        self.data_table.setColumnCount(1 + len(days_of_week))
        headers = ["Имя"] + days_of_week
        self.data_table.setHorizontalHeaderLabels(headers)

        for row, (name, schedule) in enumerate(schedule_data.items()):
            self.data_table.insertRow(row)
            self.data_table.setItem(row, 0, QTableWidgetItem(name))

            for col, day in enumerate(days_of_week, start=1):
                day_schedule = schedule.schedule_for_days.get(day, [False] * num_pairs)

                cell_widget = QWidget()
                cell_layout = QHBoxLayout(cell_widget)
                cell_layout.setContentsMargins(0, 0, 0, 0)

                # Создаем по одному QCheckBox для каждой пары в день
                for slot in day_schedule:
                    check_box = QCheckBox()
                    check_box.setChecked(slot)
                    cell_layout.addWidget(check_box)

                cell_widget.setLayout(cell_layout)
                self.data_table.setCellWidget(row, col, cell_widget)

    def _save_schedule_changes(self, schedule_data):
        for row in range(self.data_table.rowCount()):
            name = self.data_table.item(row, 0).text()

            updated_schedule = {}
            for col in range(1, self.data_table.columnCount()):
                day = self.data_table.horizontalHeaderItem(col).text()

                cell_widget = self.data_table.cellWidget(row, col)
                day_schedule = [check_box.isChecked() for check_box in cell_widget.findChildren(QCheckBox)]

                updated_schedule[day] = day_schedule

            schedule_data[name].schedule_for_days = updated_schedule

    def save_changes(self):
        if self.current_variable is None:
            return

        invalid_data = []
        table_name = self.current_variable
        ru_table_name = next((key for key, value in self.vars_to_redact.items() if value == table_name), None)

        try:
            variable_data = getattr(data, table_name)

            if table_name == "groups_shift":
                for row in range(self.data_table.rowCount()):
                    try:
                        group = self.data_table.item(row, 0).text()
                        pair = int(self.data_table.item(row, 1).text())
                        time_str = self.data_table.item(row, 2).text()
                        start_time = time.fromisoformat(time_str.split(" - ")[0])
                        end_time = time.fromisoformat(time_str.split(" - ")[1])
                        pair_type = self.data_table.item(row, 3).text()
                        variable_data[group][pair] = db.PairTime(start_time, end_time, pair_type)
                    except (ValueError, AttributeError, IndexError):
                        row_data = [self.data_table.item(row, col).text() for col in
                                    range(self.data_table.columnCount())]
                        invalid_data.append(row_data)

            elif table_name == "discipline_hours":
                for row in range(self.data_table.rowCount()):
                    try:
                        group = self.data_table.item(row, 0).text()
                        discipline = self.data_table.item(row, 1).text()
                        hours = int(self.data_table.item(row, 2).text())
                        variable_data[group][discipline] = hours
                    except (ValueError, AttributeError):
                        row_data = [self.data_table.item(row, col).text() for col in
                                    range(self.data_table.columnCount())]
                        invalid_data.append(row_data)

            elif table_name == "teachers":
                for row in range(self.data_table.rowCount()):
                    try:
                        name = self.data_table.item(row, 0).text()
                        disciplines = set(self.data_table.item(row, 1).text().split(", "))
                        groups = set(self.data_table.item(row, 2).text().split(", "))
                        variable_data[name] = db.Teacher(name, disciplines, groups)
                    except AttributeError:
                        row_data = [self.data_table.item(row, col).text() for col in
                                    range(self.data_table.columnCount())]
                        invalid_data.append(row_data)

            elif table_name == "rooms":
                for row in range(self.data_table.rowCount()):
                    try:
                        room = self.data_table.item(row, 0).text()
                        is_online = self.data_table.item(row, 1).text() == "Да"
                        variable_data[room] = db.Room(is_online)
                    except AttributeError:
                        row_data = [self.data_table.item(row, col).text() for col in
                                    range(self.data_table.columnCount())]
                        invalid_data.append(row_data)

            elif table_name in ["teachers_work_hours", "rooms_availability_hours"]:
                try:
                    self._save_schedule_changes(variable_data)
                except Exception as e:
                    QMessageBox.warning(self, "Ошибка", f"Ошибка при сохранении расписания: {str(e)}")
                    return

            setattr(data, table_name, variable_data)

        except Exception as e:
            QMessageBox.critical(self, "Критическая ошибка", f"Не удалось сохранить изменения: {str(e)}")
            return

        if invalid_data:
            # Подготовка информации о первых трех ошибках
            error_message = f"Обнаружены ошибки в таблице '{ru_table_name}'. Проверьте следующие строки:\n"
            for i, row_data in enumerate(invalid_data[:3]):
                error_message += f"\nСтрока {i + 1}: " + ", ".join(row_data)

            QMessageBox.warning(self, "Ошибка данных", error_message)
        else:
            QMessageBox.information(self, "Сохранение", f"Данные в таблице '{ru_table_name}' успешно сохранены.")


class ErrorDialog(QDialog):
    def __init__(self, errors: list, remaining_data: db.Data):
        super().__init__()
        self.setWindowTitle("Ошибки при генерации расписания")
        self.setGeometry(100, 100, 800, 500)  # Увеличиваем ширину окна для списка оставшихся часов
        self.setWindowFlags(self.windowFlags() | Qt.WindowContextHelpButtonHint)  # Добавляем кнопку справки

        # Загрузка стиля
        with open("ErrorDialog.css", "r") as f:
            self.setStyleSheet(f.read())

        # Основной макет с горизонтальным расположением для левой и правой частей
        main_layout = QHBoxLayout(self)

        # Левая часть (список ошибок и информация о текущей ошибке)
        left_layout = QVBoxLayout()
        self.errors = errors

        # Список ошибок
        self.error_list_widget = QListWidget()
        self.error_list_widget.addItems(
            [f"{i + 1} / Группа: {error.group}, Дисциплина: {error.discipline}" for i, error in enumerate(errors)]
        )
        self.error_list_widget.currentItemChanged.connect(self.display_error_info)

        # Заголовок и описание ошибок
        self.help_text = QLabel("Невозможно поставить пары для данных дисциплин и групп:")
        left_layout.addWidget(self.help_text)
        left_layout.addWidget(self.error_list_widget)

        self.group_label = QLabel("| Группа: ")
        self.discipline_label = QLabel("| Дисциплина: ")
        self.hours_label = QLabel("| Оставшиеся часы: ")
        left_layout.addWidget(self.group_label)
        left_layout.addWidget(self.discipline_label)
        left_layout.addWidget(self.hours_label)

        main_layout.addLayout(left_layout)

        # Правая часть (список дисциплин с оставшимися часами)
        right_layout = QVBoxLayout()
        self.remaining_hours_list_widget = QListWidget()

        # Заголовок для оставшихся часов
        remaining_hours_label = QLabel("Дисциплины с оставшимися часами:")
        right_layout.addWidget(remaining_hours_label)
        right_layout.addWidget(self.remaining_hours_list_widget)

        # Добавление данных об оставшихся часах в виджет
        for group, disciplines in remaining_data.discipline_hours.items():
            for discipline, hours in disciplines.items():
                if hours > 0:
                    self.remaining_hours_list_widget.addItem(
                        f"Группа: {group},\nДисциплина: {discipline},\nОсталось часов: {hours}")

        self.back_button = QPushButton("Назад")
        self.back_button.clicked.connect(self.close)
        right_layout.addWidget(self.back_button)

        main_layout.addLayout(right_layout)

    def display_error_info(self, current):
        def pair_text(count):
            if count % 10 == 1: return "пара"
            if count % 10 in [2, 3, 4]: return "пары"
            return "пар"

        if current:
            current_error = self.errors[int(current.text().split("/")[0].strip()) - 1]
            self.group_label.setText(f"| Группа: {current_error.group}")
            self.discipline_label.setText(f"| Дисциплина: {current_error.discipline}")
            self.hours_label.setText(
                f"| Оставшиеся часы: {current_error.hours} (= {current_error.hours // 2} {pair_text(current_error.hours // 2)})")

    def event(self, event):
        if event.type() == QEvent.Type(124):
            self.show_help()
            return True
        return super().event(event)

    def show_help(self):
        # Отображение справки
        help_message = (
            "В этом окне отображаются ошибки, возникшие при генерации расписания.\n\n"
            "Список вверху содержит информацию о группе и дисциплине, для которой не удалось "
            "поставить пары. Выберите элемент из списка, чтобы увидеть подробные данные "
            "о выбранной ошибке ниже.\n\n"
            "С правой стороны отображается список дисциплин с оставшимися часами, "
            "которые ещё не были распределены в расписании."
        )
        QMessageBox.information(self, "Справка", help_message)


class MainWindow(QMainWindow):
    def __init__(self):
        self.current_schedule: dict[str, list[db.Pair]] = None
        self.remaining_data = None
        self.errors = []

        super().__init__()

        self.setWindowTitle("Составление расписания")
        self.setWindowIcon(QIcon("icon.png"))
        self.setGeometry(100, 100, 1400, 800)

        self.central_widget = QWidget()
        self.setCentralWidget(self.central_widget)

        self.layout = QHBoxLayout(self.central_widget)

        self.table_widget = QTableWidget(1, 1)
        self.table_widget.setItem(0, 0, QTableWidgetItem("Здесь будет созданное расписание"))
        self.table_widget.setEditTriggers(QAbstractItemView.NoEditTriggers)

        self.button_layout = QVBoxLayout()

        self.generate_schedule_button = QPushButton("Сгенерировать расписание")
        self.generate_schedule_button.clicked.connect(self.generate_schedule)
        self.button_layout.addWidget(self.generate_schedule_button)

        self.error_button = QPushButton(f"Ошибки: {len(self.errors)}")
        self.error_button.setEnabled(False)
        self.error_button.clicked.connect(self.show_errors)
        self.button_layout.addWidget(self.error_button)

        self.export_schedule_button = QPushButton("Выгрузить расписание")
        self.export_schedule_button.clicked.connect(self.export_schedule)
        self.button_layout.addWidget(self.export_schedule_button)

        self.input_data_button = QPushButton("Ввод данных")
        self.input_data_button.clicked.connect(self.input_data)
        self.button_layout.addWidget(self.input_data_button)

        self.layout.addWidget(self.table_widget)
        self.layout.addLayout(self.button_layout)

        self.button_layout.addStretch()

        with open("MainWindow.css", "r") as f:
            self.setStyleSheet(f.read())
        self.resize_columns()

    def resize_columns(self):
        self.table_widget.resizeColumnsToContents()

    def generate_schedule(self):
        sch = schedule_maker.make_full_schedule(data)
        self.current_schedule = sch
        self.errors = sch.errors
        self.remaining_data = sch.remaining_data
        headers = ["Группа", "День", "Время", "Форма", "Предмет", "Педагог", "Каб."]
        self.table_widget.setRowCount(0)
        self.table_widget.setColumnCount(len(headers))
        self.table_widget.setHorizontalHeaderLabels(headers)

        rows: list[str] = []
        for group, pairs in sch.pairs.items():
            for pair in pairs:
                rows.append([group, pair.day, pair.pair_time.get_str(), pair.pair_type, pair.discipline, pair.teacher,
                             pair.classroom])

        for row in rows:
            current_row = self.table_widget.rowCount()
            self.table_widget.insertRow(current_row)
            for column in range(len(row)):
                self.table_widget.setItem(current_row, column, QTableWidgetItem(row[column]))

        self.resize_columns()

        self.error_button.setText(f"Ошибки: {len(self.errors)}")
        self.error_button.setEnabled(len(self.errors) > 0)

    def show_errors(self):
        if self.errors:
            error_dialog = ErrorDialog(self.errors, self.remaining_data)
            error_dialog.exec_()

    def export_schedule(self):
        try:
            row_count = self.table_widget.rowCount()
            column_count = self.table_widget.columnCount()

            exp_data = []
            for row in range(row_count):
                row_data = []
                for column in range(column_count):
                    item = self.table_widget.item(row, column)
                    row_data.append(item.text() if item else "")
                exp_data.append(row_data)

            headers = [self.table_widget.horizontalHeaderItem(i).text() for i in range(column_count)]

            workbook = Workbook()
            sheet = workbook.active

            for col_num, header in enumerate(headers, 1):
                sheet.cell(row=1, column=col_num, value=header)

            for row_num, row_data in enumerate(exp_data, 2):
                for col_num, value in enumerate(row_data, 1):
                    sheet.cell(row=row_num, column=col_num, value=value)

            for col in range(1, column_count + 1):
                max_length = 0
                column_letter = get_column_letter(col)
                for row in range(1, row_count + 2):
                    cell_value = sheet[f"{column_letter}{row}"].value
                    if cell_value:
                        max_length = max(max_length, len(str(cell_value)))
                adjusted_width = (max_length + 2)
                sheet.column_dimensions[column_letter].width = adjusted_width

            workbook.save("расписание.xlsx")
            resp = QMessageBox.information(
                self,
                "Успех",
                "Расписание успешно выгружено\nОткрыть xlsx файл?",
                buttons=QMessageBox.Ok | QMessageBox.Cancel
            )

            if resp == QMessageBox.Ok:
                import os
                os.startfile("расписание.xlsx")
        except AttributeError:
            QMessageBox.warning(self, "Ошибка", "Сгенерируйте расписание перед выгрузкой")
        except Exception as e:
            QMessageBox.warning(self, "Ошибка", str(e))
            raise e

    @staticmethod
    def input_data():
        input_dialog = InputDataDialog()
        input_dialog.exec_()


def global_exception_handler(exctype, value, traceback):
    print("Произошла необработанная ошибка:", value)
    with open("error_log.txt", "a") as f:
        f.write(f"Произошла не обработанная ошибка: {value}\n")
    db.save_data(data)
    sys.__excepthook__(exctype, value, traceback)
    sys.exit(1)


sys.excepthook = global_exception_handler

if __name__ == "__main__":
    data = db.get_data()
    app = QApplication(sys.argv)
    window = MainWindow()
    window.show()
    ex_code = app.exec_()
    db.save_data(data)
    sys.exit(ex_code)
